"""SlideRenderer — DeckPlan -> python-pptx Presentation -> .pptx on disk."""

from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from ..config.loader import DeckConfig
from ..config.themes import resolve_theme
from ..core.plan import DeckPlan, Slide
from .layout_engine import apply_layout

log = logging.getLogger(__name__)


class SlideRenderer:
    def __init__(self, config: DeckConfig) -> None:
        self.config = config

    def render(self, plan: DeckPlan, out_path: Path) -> Path:
        theme, template_path = resolve_theme(plan.theme or self.config.theme)
        if template_path.exists():
            prs = Presentation(str(template_path))
        else:
            log.warning("theme template %s missing; using blank presentation", template_path)
            prs = Presentation()

        self._render_title(prs, plan, theme)

        for slide in plan.slides:
            layout = apply_layout(prs, slide, theme)
            self._fill_body(layout, slide, theme)
            if slide.speaker_notes:
                layout.notes_slide.notes_text_frame.text = slide.speaker_notes

        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
        return out_path

    def _render_title(self, prs: Presentation, plan: DeckPlan, theme) -> None:
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = plan.title
        if len(slide.placeholders) > 1 and plan.subtitle:
            slide.placeholders[1].text = plan.subtitle

    def _fill_body(self, slide, plan_slide: Slide, theme) -> None:
        body = None
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                body = ph
                break
        if body is None:
            return
        tf = body.text_frame
        tf.clear()
        for i, block in enumerate(plan_slide.blocks):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = block.text
            p.font.size = Pt(18 if block.kind == "bullet" else 16)
            if block.kind == "bullet":
                p.level = 0